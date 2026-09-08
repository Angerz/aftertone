from __future__ import annotations

import os
import shutil
import tempfile
from io import BytesIO
from decimal import Decimal
from pathlib import Path

os.environ["AFTERTONE_DATABASE_URL"] = f"sqlite:///{Path(__file__).parent / 'aftertone-api-test.sqlite3'}"
TEST_COVER_DIR = Path(tempfile.mkdtemp(prefix="aftertone-cover-test-"))
os.environ["AFTERTONE_COVER_DIR"] = str(TEST_COVER_DIR)

import pytest
import httpx
from sqlalchemy import func, select
from fastapi import HTTPException, UploadFile
from fastapi.testclient import TestClient
from PIL import Image
from pydantic import ValidationError
from starlette.datastructures import Headers

from app.config import cover_dir
from app.database import Base, SessionLocal, engine
from app.imports.legacy_excel import _text, commit_rows, preview_workbook
from app.imports.pre_formula import PreFormulaError, legacy_adjustment_value, parse_pre_formula
from app.main import album_facets, app, clear_revisit_mark, create_album, create_revision, delete_cover, get_album, get_artist, get_revision, import_cover_from_url, list_albums, list_artists, list_revisions, mark_for_revisit, update_album, update_track_artists, upload_cover
from app.ratings.calculator import TrackScore, calculate_rating
from app.models.music import Album, LegacyRating, Track
from app.schemas.music import AlbumCreate, AlbumUpdate, CoverFromUrl, LegacyReconciliationRequest, LegacyTrackMapping, RatingRevisionCreate, RevisitUpdate, TrackCreditsUpdate, TrackUpdate
from app.services.legacy_reconciliation import ReconciliationError, reconcile_legacy_rating
from app.services.covers import CoverError, MAX_COVER_BYTES, download_cover_from_url, save_cover_data


@pytest.fixture(autouse=True)
def database() -> None:
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    shutil.rmtree(TEST_COVER_DIR, ignore_errors=True)
    TEST_COVER_DIR.mkdir()
    yield
    Base.metadata.drop_all(engine)
    shutil.rmtree(TEST_COVER_DIR, ignore_errors=True)


def album_payload(title: str = "Grace") -> dict[str, object]:
    return {
        "title": title, "artists": [{"name": "Jeff Buckley"}], "year": 1994, "release_type": "album",
        "tracks": [{"position": 2, "title": "Grace"}, {"position": 1, "title": "Mojo Pin"}],
    }


def create_test_album(title: str = "Grace"):
    with SessionLocal() as session:
        return create_album(AlbumCreate.model_validate(album_payload(title)), session)


def revision_payload(album) -> dict[str, object]:
    return {
        "coherence": "7", "coherence_notes": "Cohesive sequencing.", "emotion": "10",
        "emotion_notes": "Overwhelming.", "album_notes": "A complete note.",
        "tracks": [
            {"track_id": album.tracks[0].id, "score": "10", "include_in_pre_rating": True, "notes": "Opener."},
            {"track_id": album.tracks[1].id, "score": "9", "include_in_pre_rating": False, "notes": "Title track."},
        ],
    }


def image_upload(image_format: str = "PNG", content_type: str = "image/png") -> UploadFile:
    data = BytesIO()
    Image.new("RGB", (1600, 800), color="navy").save(data, format=image_format)
    data.seek(0)
    return UploadFile(file=data, filename=f"cover.{image_format.lower()}", headers=Headers({"content-type": content_type}))


def image_bytes(image_format: str) -> bytes:
    data = BytesIO()
    Image.new("RGB", (40, 30), color="navy").save(data, format=image_format)
    return data.getvalue()


def public_dns(monkeypatch) -> None:
    monkeypatch.setattr("app.services.covers.socket.getaddrinfo", lambda *_args, **_kwargs: [(None, None, None, None, ("93.184.216.34", 443))])


def test_create_get_and_list_albums_with_ordered_tracks() -> None:
    created = create_test_album()
    assert [track.position for track in created.tracks] == [1, 2]
    assert created.year == 1994
    with SessionLocal() as session:
        found = get_album(created.id, session)
        assert [track.title for track in found.tracks] == ["Mojo Pin", "Grace"]
        listed = list_albums(session=session)
        assert [album.id for album in listed.items] == [created.id]
        assert listed.items[0].latest_revision is None


def test_update_album_metadata_rejects_other_album_state() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        updated = update_album(album.id, AlbumUpdate(title="Grace (Remastered)", artists=[{"name": "Jeff Buckley"}], year=1995, release_type="ep"), session)
    assert (updated.title, [artist.name for artist in updated.artists], updated.year, updated.release_type.value) == ("Grace (Remastered)", ["Jeff Buckley"], 1995, "ep")
    with SessionLocal() as session, pytest.raises(HTTPException) as missing:
        update_album(99999, AlbumUpdate(title="Missing", artists=[{"name": "Nobody"}], year=None, release_type="album"), session)
    assert missing.value.status_code == 404
    with pytest.raises(ValidationError):
        AlbumUpdate.model_validate({"title": "Grace", "artists": [{"name": "Jeff Buckley"}], "year": 1994, "release_type": "album", "cover_filename": "not-allowed.webp"})
    with pytest.raises(ValidationError):
        AlbumUpdate.model_validate({"title": " ", "artists": [{"name": "Jeff Buckley"}], "year": 999, "release_type": "album"})


def test_update_album_tracklist_reorders_renames_and_preserves_revision_snapshots() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        revision = create_revision(album.id, RatingRevisionCreate.model_validate(revision_payload(album)), session)
        snapshots = [(item.track_id, item.title, item.position) for item in revision.tracks]
        updated = update_album(album.id, AlbumUpdate(title=album.title, artists=[{"name": "Jeff Buckley"}], year=1994, release_type="album", tracks=[TrackUpdate(id=album.tracks[1].id, title="Grace renamed"), TrackUpdate(id=album.tracks[0].id, title="Mojo Pin renamed"), TrackUpdate(title="New track")]), session)
        stored_revision = get_revision(album.id, revision.id, session)
    assert [(track.title, track.position) for track in updated.tracks] == [("Grace renamed", 1), ("Mojo Pin renamed", 2), ("New track", 3)]
    assert [(item.track_id, item.title, item.position) for item in stored_revision.tracks] == snapshots


def test_album_list_includes_only_the_latest_revision_summary() -> None:
    album = create_test_album()
    first_payload = RatingRevisionCreate.model_validate(revision_payload(album))
    with SessionLocal() as session:
        first = create_revision(album.id, first_payload, session)
    second_data = revision_payload(album)
    second_data["emotion"] = "5"
    with SessionLocal() as session:
        second = create_revision(album.id, RatingRevisionCreate.model_validate(second_data), session)
        latest = list_albums(session=session).items[0].latest_revision
    assert latest is not None
    assert latest.id == second.id
    assert latest.id != first.id
    assert latest.pre_rating == second.pre_rating
    assert latest.final_rating == second.final_rating


def test_album_list_serializes_rating_decimals_as_json_numbers() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        create_revision(album.id, RatingRevisionCreate.model_validate(revision_payload(album)), session)
        session.add(LegacyRating(
            album_id=album.id, coherence=Decimal("7"), emotion=Decimal("8"),
            extracted_scores=["10", "9"], pre_formula="=(10+9)/2",
            legacy_final_rating=Decimal("8.5"), computed_final_rating=Decimal("8.75"),
        ))
        session.commit()
    with SessionLocal() as session:
        payload = list_albums(session=session).items[0].model_dump(mode="json")
    assert isinstance(payload["latest_revision"]["pre_rating"], float)
    assert isinstance(payload["latest_revision"]["final_rating"], float)
    assert isinstance(payload["latest_legacy_rating"]["legacy_final_rating"], float)
    assert isinstance(payload["latest_legacy_rating"]["computed_final_rating"], float)


def test_album_artists_and_track_featured_credits_are_structured() -> None:
    payload = album_payload()
    payload["artists"] = [{"name": "Primary Artist"}, {"name": "Second Artist"}]
    with SessionLocal() as session:
        created = create_album(AlbumCreate.model_validate(payload), session)
        assert [artist.name for artist in created.artists] == ["Primary Artist", "Second Artist"]
        inherited = get_album(created.id, session).tracks[0]
        assert inherited.uses_album_artists is True
        assert [artist.name for artist in inherited.primary_artists] == ["Primary Artist", "Second Artist"]
        featured = update_track_artists(inherited.id, TrackCreditsUpdate(featured_artist_ids=[created.artists[1].id]), session)
        assert [artist.name for artist in featured.featured_artists] == ["Second Artist"]
        detail = get_artist(created.artists[1].id, session)
    assert [album.title for album in detail.albums] == ["Grace"]
    assert detail.albums[0].rating is None
    assert [appearance.track_title for appearance in detail.featured_appearances] == ["Mojo Pin"]


def test_artist_projects_include_the_current_effective_rating() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        revision = create_revision(album.id, RatingRevisionCreate.model_validate(revision_payload(album)), session)
        detail = get_artist(album.artists[0].id, session)
    assert detail.albums[0].rating == revision.final_rating


def test_paginated_library_filters_searches_and_facets_before_slicing() -> None:
    rows = [("Nineties", 2019, ["Alpha"]), ("First", 2020, ["Alpha", "Guest"]), ("Second", 2024, ["Beta"]), ("Last", 2029, ["Gamma"]), ("Future", 2030, ["Delta"]), ("Unknown", None, ["Unknown"])]
    with SessionLocal() as session:
        for title, year, artists in rows:
            create_album(AlbumCreate.model_validate({"title": title, "artists": [{"name": artist} for artist in artists], "year": year, "release_type": "album", "tracks": [{"position": 1, "title": "Track"}]}), session)
        first = list_albums(page=1, page_size=2, sort="title", session=session)
        second = list_albums(page=2, page_size=2, sort="title", session=session)
        twenty_twenty = list_albums(decade=2020, page=1, page_size=10, sort="year", session=session)
        artist_search = list_albums(search="guest", session=session)
        empty = list_albums(page=99, page_size=2, session=session)
        facets = album_facets(session)
    assert first.total == 6 and first.total_pages == 3 and [item.title for item in first.items] == ["First", "Future"]
    assert [item.title for item in second.items] == ["Last", "Nineties"]
    assert [item.title for item in twenty_twenty.items] == ["Last", "Second", "First"]
    assert [item.title for item in artist_search.items] == ["First"]
    assert empty.items == [] and empty.page == 99
    assert facets.decades == {2010: [2019], 2020: [2020, 2024, 2029], 2030: [2030]}


def test_artist_pagination_search_and_ordering() -> None:
    with SessionLocal() as session:
        for name in ["Zulu", "Alpha", "Beta"]:
            create_album(AlbumCreate.model_validate({"title": f"{name} album", "artists": [{"name": name}], "year": 2020, "release_type": "album", "tracks": [{"position": 1, "title": "Track"}]}), session)
        first = list_artists(page=1, page_size=2, session=session)
        searched = list_artists(search="et", session=session)
        empty = list_artists(page=9, page_size=2, session=session)
    assert [item.name for item in first.items] == ["Alpha", "Beta"]
    assert first.total == 3 and first.total_pages == 2
    assert [item.name for item in searched.items] == ["Beta"]
    assert empty.items == [] and empty.total == 3


def test_create_and_read_complete_calculated_revision() -> None:
    album = create_test_album()
    payload = RatingRevisionCreate.model_validate(revision_payload(album))
    with SessionLocal() as session:
        revision = create_revision(album.id, payload, session)
        expected = calculate_rating([TrackScore(Decimal("10"), True), TrackScore(Decimal("9"), False)],
                                    coherence=Decimal("7"), emotion=Decimal("10"))
        assert revision.pre_rating == expected.pre_rating
        assert revision.bad_experience == expected.bad_experience
        assert revision.final_rating == expected.final_rating
        assert revision.album_notes == "A complete note."
        assert revision.coherence_notes == "Cohesive sequencing."
        assert revision.emotion_notes == "Overwhelming."
        assert revision.tracks[0].notes == "Opener."
        assert revision.tracks[1].include_in_pre_rating is False
    with SessionLocal() as session:
        fetched = get_revision(album.id, revision.id, session)
        assert fetched == revision
        assert list_revisions(album.id, session)[0].id == revision.id

    no_pre_payload = revision_payload(album)
    for track in no_pre_payload["tracks"]:
        track["include_in_pre_rating"] = False
    with SessionLocal() as session:
        no_pre_revision = create_revision(album.id, RatingRevisionCreate.model_validate(no_pre_payload), session)
        assert no_pre_revision.id != revision.id
        assert no_pre_revision.pre_rating is None
        assert no_pre_revision.bad_experience is None
        assert no_pre_revision.final_rating is None
        assert len(list_revisions(album.id, session)) == 2


@pytest.mark.parametrize(("field", "value"), [("coherence", "10.1"), ("coherence", "-0.1"),
                                                   ("emotion", "10.1"), ("emotion", "-0.1")])
def test_revision_rejects_out_of_range_album_scores(field: str, value: str) -> None:
    album = create_test_album()
    payload = revision_payload(album)
    payload[field] = value
    with pytest.raises(ValidationError):
        RatingRevisionCreate.model_validate(payload)


def test_revision_rejects_client_supplied_calculated_value() -> None:
    payload = revision_payload(create_test_album())
    payload["bad_experience"] = "1"
    with pytest.raises(ValidationError):
        RatingRevisionCreate.model_validate(payload)


def test_revision_validates_complete_unique_album_tracklist() -> None:
    album = create_test_album()
    other_album = create_test_album("Sketches")
    payload = revision_payload(album)
    payload["tracks"][0]["score"] = "10.1"
    with pytest.raises(ValidationError):
        RatingRevisionCreate.model_validate(payload)
    payload["tracks"][0]["score"] = "10"
    payload["tracks"].append(payload["tracks"][0].copy())
    with SessionLocal() as session, pytest.raises(HTTPException) as duplicate:
        create_revision(album.id, RatingRevisionCreate.model_validate(payload), session)
    assert duplicate.value.status_code == 422
    payload["tracks"].pop()
    payload["tracks"][1]["track_id"] = 99999
    with SessionLocal() as session, pytest.raises(HTTPException) as nonexistent:
        create_revision(album.id, RatingRevisionCreate.model_validate(payload), session)
    assert nonexistent.value.status_code == 422
    payload["tracks"][1]["track_id"] = other_album.tracks[0].id
    with SessionLocal() as session, pytest.raises(HTTPException) as foreign:
        create_revision(album.id, RatingRevisionCreate.model_validate(payload), session)
    assert foreign.value.status_code == 422
    payload["tracks"] = payload["tracks"][:1]
    with SessionLocal() as session, pytest.raises(HTTPException) as incomplete:
        create_revision(album.id, RatingRevisionCreate.model_validate(payload), session)
    assert incomplete.value.status_code == 422


def test_revision_and_album_not_found_or_mismatched_path() -> None:
    with SessionLocal() as session, pytest.raises(HTTPException) as missing_album:
        get_album(404, session)
    assert missing_album.value.status_code == 404
    album = create_test_album()
    with SessionLocal() as session:
        revision = create_revision(album.id, RatingRevisionCreate.model_validate(revision_payload(album)), session)
    other_album = create_test_album("Sketches")
    with SessionLocal() as session, pytest.raises(HTTPException) as wrong_album:
        get_revision(other_album.id, revision.id, session)
    assert wrong_album.value.status_code == 404


@pytest.mark.parametrize(("image_format", "content_type"), [("JPEG", "image/jpeg"), ("PNG", "image/png"), ("WEBP", "image/webp")])
def test_upload_cover_normalizes_supported_images(image_format: str, content_type: str) -> None:
    album = create_test_album()
    with SessionLocal() as session:
        response = upload_cover(album.id, image_upload(image_format, content_type), session)
    assert response.cover_url is not None
    filename = response.cover_url.rsplit("/", 1)[-1]
    assert filename.endswith(".webp")
    assert (cover_dir() / filename).is_file()
    with Image.open(cover_dir() / filename) as stored:
        assert stored.format == "WEBP"
        assert max(stored.size) <= 1200


def test_cover_upload_rejects_invalid_content_and_missing_album() -> None:
    album = create_test_album()
    invalid_type = UploadFile(file=BytesIO(b"not an image"), filename="cover.txt", headers=Headers({"content-type": "text/plain"}))
    with SessionLocal() as session, pytest.raises(HTTPException) as invalid:
        upload_cover(album.id, invalid_type, session)
    assert invalid.value.status_code == 422
    corrupt_image = UploadFile(file=BytesIO(b"not an image"), filename="cover.png", headers=Headers({"content-type": "image/png"}))
    with SessionLocal() as session, pytest.raises(HTTPException) as corrupt:
        upload_cover(album.id, corrupt_image, session)
    assert corrupt.value.status_code == 422
    with SessionLocal() as session, pytest.raises(HTTPException) as missing:
        upload_cover(99999, image_upload(), session)
    assert missing.value.status_code == 404


def test_cover_upload_rejects_oversized_file() -> None:
    album = create_test_album()
    oversized = UploadFile(
        file=BytesIO(b"x" * (10 * 1024 * 1024 + 1)), filename="cover.png", headers=Headers({"content-type": "image/png"})
    )
    with SessionLocal() as session, pytest.raises(HTTPException) as error:
        upload_cover(album.id, oversized, session)
    assert error.value.status_code == 413


@pytest.mark.parametrize(("image_format", "content_type"), [("JPEG", "image/jpeg"), ("PNG", "image/png"), ("WEBP", "image/webp")])
def test_download_cover_from_public_url_uses_local_cover_pipeline(monkeypatch, image_format: str, content_type: str) -> None:
    public_dns(monkeypatch)
    client = httpx.Client(transport=httpx.MockTransport(lambda _request: httpx.Response(200, headers={"content-type": content_type}, content=image_bytes(image_format))))
    try:
        filename = download_cover_from_url("https://covers.example/album", client)
    finally:
        client.close()
    assert filename.endswith(".webp")
    with Image.open(cover_dir() / filename) as stored:
        assert stored.format == "WEBP"


def test_download_cover_url_rejects_unsafe_and_invalid_responses(monkeypatch) -> None:
    with pytest.raises(CoverError, match="http or https"):
        download_cover_from_url("file:///tmp/cover.jpg")
    monkeypatch.setattr("app.services.covers.socket.getaddrinfo", lambda *_args, **_kwargs: [(None, None, None, None, ("127.0.0.1", 80))])
    with pytest.raises(CoverError, match="local or private"):
        download_cover_from_url("http://localhost/cover.jpg")
    monkeypatch.setattr("app.services.covers.socket.getaddrinfo", lambda *_args, **_kwargs: [(None, None, None, None, ("10.0.0.5", 80))])
    with pytest.raises(CoverError, match="local or private"):
        download_cover_from_url("http://10.0.0.5/cover.jpg")
    public_dns(monkeypatch)
    cases = [
        httpx.Response(200, headers={"content-type": "text/html"}, content=b"not an image"),
        httpx.Response(200, headers={"content-type": "image/jpeg"}, content=b"corrupt"),
        httpx.Response(200, headers={"content-type": "image/jpeg", "content-length": str(MAX_COVER_BYTES + 1)}),
    ]
    for response in cases:
        client = httpx.Client(transport=httpx.MockTransport(lambda _request, result=response: result))
        try:
            with pytest.raises(CoverError):
                download_cover_from_url("https://covers.example/album", client)
        finally:
            client.close()


def test_download_cover_url_revalidates_redirect_destinations(monkeypatch) -> None:
    monkeypatch.setattr(
        "app.services.covers.socket.getaddrinfo",
        lambda host, *_args, **_kwargs: [(None, None, None, None, (("127.0.0.1" if host == "localhost" else "93.184.216.34"), 80))],
    )
    client = httpx.Client(transport=httpx.MockTransport(lambda _request: httpx.Response(302, headers={"location": "http://localhost/private.jpg"})))
    try:
        with pytest.raises(CoverError, match="local or private"):
            download_cover_from_url("https://covers.example/album", client)
    finally:
        client.close()


def test_import_cover_from_url_replaces_existing_cover(monkeypatch) -> None:
    album = create_test_album()
    first_filename = save_cover_data(image_bytes("PNG"), "image/png")
    second_filename = save_cover_data(image_bytes("JPEG"), "image/jpeg")
    with SessionLocal() as session:
        stored = session.get(Album, album.id)
        assert stored is not None
        stored.cover_filename = first_filename
        session.commit()
        monkeypatch.setattr("app.main.download_cover_from_url", lambda _url: second_filename)
        response = import_cover_from_url(album.id, CoverFromUrl(url="https://covers.example/replacement.jpg"), session)
    assert response.cover_url is not None
    assert not (cover_dir() / first_filename).exists()
    assert (cover_dir() / second_filename).is_file()


def test_replace_and_delete_cover_removes_old_file() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        first = upload_cover(album.id, image_upload(), session)
        first_filename = first.cover_url.rsplit("/", 1)[-1] if first.cover_url else ""
    with SessionLocal() as session:
        second = upload_cover(album.id, image_upload("JPEG", "image/jpeg"), session)
        second_filename = second.cover_url.rsplit("/", 1)[-1] if second.cover_url else ""
    assert not (cover_dir() / first_filename).exists()
    assert (cover_dir() / second_filename).is_file()
    with SessionLocal() as session:
        deleted = delete_cover(album.id, session)
    assert deleted.cover_url is None
    assert not (cover_dir() / second_filename).exists()
    with SessionLocal() as session:
        assert delete_cover(album.id, session).cover_url is None


def test_legacy_pre_formula_parses_decimal_comma_and_audits_grace() -> None:
    scores, denominator, suffix = parse_pre_formula("=(10+10+10+10+10+10+10+10+9,5+10+10)/11")
    assert denominator == 11
    assert scores[8] == Decimal("9.5")
    assert len(scores) == 11
    assert suffix == ""
    zero_scores, zero_denominator, zero_suffix = parse_pre_formula("=(10+9.5+8)/3+0-0")
    assert (zero_scores, zero_denominator, zero_suffix) == ([Decimal("10"), Decimal("9.5"), Decimal("8")], 3, "+0-0")
    positive_scores, positive_denominator, positive_suffix = parse_pre_formula("=(10+9,5+8)/3+0,1-0")
    assert (positive_scores, positive_denominator, positive_suffix) == ([Decimal("10"), Decimal("9.5"), Decimal("8")], 3, "+0,1-0")
    assert legacy_adjustment_value(positive_suffix) == Decimal("0.1")
    with pytest.raises(PreFormulaError):
        parse_pre_formula("=AVERAGE(10,9)")
    with pytest.raises(PreFormulaError):
        parse_pre_formula("=(10+9+8)/3+A1")


def test_legacy_preview_marks_audit_discrepancies_as_warnings() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "AÑO", "CALIFICACIÓN", None, "PRE-CALIFICACIÓN", "COHERENCIA", "MALA EXP", "EMOCIÓN"])
    sheet.append(["Grace", "Jeff Buckley", 1994, "=F2+0.1", None, "=(10+4)/2", 7, 0, 10])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        preview = preview_workbook(data.getvalue(), session)
        imported, skipped, failed = commit_rows(preview, {2}, session)
        legacy = session.scalar(select(LegacyRating))
    assert preview[0].status == "warning"
    assert (imported, skipped, failed) == (1, 0, 0)
    assert legacy is not None
    assert preview[0].legacy_final_rating is None
    assert legacy.legacy_final_rating is None
    assert preview[0].computed_final_rating == calculate_rating(
        [TrackScore(Decimal("10")), TrackScore(Decimal("4"))], coherence=Decimal("7"), emotion=Decimal("10")
    ).final_rating
    assert any("Legacy bad experience differs" in warning for warning in preview[0].warnings)


def test_legacy_preview_accepts_post_average_adjustments_without_using_them() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "PRE-CALIFICACIÓN", "COHERENCIA", "EMOCIÓN"])
    sheet.append(["Comfort y Música Para Volar [Live]", "Soda Stereo", "=(10+9.6+8.3+10+10+9.5+10+8.1+8.7+8.4+8.1)/11+0-0", 7, 10])
    sheet.append(["Soviet Kitsch", "Regina Spektor", "=(8.3+7.5+9.2+9.9+8.7+8.9+8.8+8.9+9.1+8.6)/10+0.1-0", 7, 10])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        preview = preview_workbook(data.getvalue(), session)
    neutral, adjusted = preview
    assert len(neutral.extracted_scores) == 11
    assert neutral.computed_pre_rating == sum(neutral.extracted_scores) / Decimal("11")
    assert not any("legacy adjustment" in warning for warning in neutral.warnings)
    assert len(adjusted.extracted_scores) == 10
    assert adjusted.computed_pre_rating == sum(adjusted.extracted_scores) / Decimal("10")
    assert adjusted.status == "warning"
    assert any("manual adjustment (+0.1)" in warning for warning in adjusted.warnings)


def test_legacy_preview_uses_scores_when_pre_denominator_does_not_match() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "PRE-CALIFICACIÓN", "COHERENCIA", "EMOCIÓN"])
    sheet.append(["SMITHEREENS", "Joji", "=(9.4+6.8+9.3+7.5+6.6+5.5+4.1+6.5+5.4)/8+0-0.25", 7, 8])
    sheet.append(["Temporada de reggaetón", "Bad Bunny", "=(7.2+1.8+4.2+8.2+8.1+6.2+5)/6", 7, 8])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        preview = preview_workbook(data.getvalue(), session)
    smithereens, temporada = preview
    assert len(smithereens.extracted_scores) == 9
    assert smithereens.status == "warning"
    assert smithereens.computed_pre_rating == sum(smithereens.extracted_scores) / Decimal("9")
    assert smithereens.computed_final_rating is not None
    assert any("denominator (8) differs from extracted score count (9)" in warning for warning in smithereens.warnings)
    assert any("manual adjustment (-0.25)" in warning for warning in smithereens.warnings)
    assert len(temporada.extracted_scores) == 7
    assert temporada.status == "warning"
    assert temporada.computed_pre_rating == sum(temporada.extracted_scores) / Decimal("7")
    assert temporada.computed_final_rating is not None
    assert any("denominator (6) differs from extracted score count (7)" in warning for warning in temporada.warnings)


def test_legacy_pre_formula_rejects_out_of_range_scores() -> None:
    with pytest.raises(PreFormulaError, match="scores must be between 0 and 10"):
        parse_pre_formula("=(11+8)/2")


def test_legacy_commit_creates_no_tracks_or_rating_revision() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "PRE-CALIFICACIÓN", "COHERENCIA", "EMOCIÓN"])
    sheet.append(["Legacy Album", "Legacy Artist", "=(10+9.5)/2", 7, 8])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        rows = preview_workbook(data.getvalue(), session)
        imported, skipped, failed = commit_rows(rows, {2}, session)
        legacy = session.scalar(select(LegacyRating))
        album = session.get(Album, legacy.album_id) if legacy else None
        track_count = len(album.tracks) if album else None
        revision_count = len(album.revisions) if album else None
    assert (imported, skipped, failed) == (1, 0, 0)
    assert legacy is not None
    assert legacy.extracted_scores == ["10", "9.5"]
    assert legacy.legacy_final_rating is None
    assert legacy.computed_final_rating == calculate_rating(
        [TrackScore(Decimal("10")), TrackScore(Decimal("9.5"))], coherence=Decimal("7"), emotion=Decimal("8")
    ).final_rating
    assert album is not None
    assert track_count == 0
    assert revision_count == 0


def test_legacy_unrated_album_creates_only_an_album_and_respects_duplicates() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "AÑO", "PRE-CALIFICACIÓN", "COHERENCIA", "MALA EXP", "EMOCIÓN"])
    sheet.append(["Entertainment", "Gang Of Four", 1979, None, None, None, None])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        rows = preview_workbook(data.getvalue(), session)
        imported, skipped, failed = commit_rows(rows, {2}, session)
        album = session.scalar(select(Album).where(Album.title == "Entertainment"))
        legacy_count = session.scalar(select(func.count(LegacyRating.id)))
        duplicate = preview_workbook(data.getvalue(), session)[0]
        duplicate_result = commit_rows([duplicate], {2}, session)
    assert rows[0].status == "unrated"
    assert rows[0].computed_final_rating is None
    assert (imported, skipped, failed) == (1, 0, 0)
    assert album is not None
    assert legacy_count == 0
    assert duplicate.status == "unrated"
    assert any("already exists" in warning for warning in duplicate.warnings)
    assert duplicate_result == (0, 1, 0)


@pytest.mark.parametrize(("title", "expected_status"), [(0, "unrated"), (0.0, "unrated"), ("0", "unrated"), (None, "error"), ("   ", "error")])
def test_legacy_title_normalization_preserves_numeric_zero(title: object, expected_status: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "PRE-CALIFICACIÓN", "COHERENCIA", "MALA EXP", "EMOCIÓN"])
    sheet.append([title, "Ichiko Aoba", None, None, None, None])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        row = preview_workbook(data.getvalue(), session)[0]
    assert row.status == expected_status
    if expected_status == "unrated":
        assert row.title == "0"
    else:
            assert any("Album and artist are required" in error for error in row.errors)


@pytest.mark.parametrize(("value", "expected"), [(0.0, "0"), (12.0, "12"), (1.5, "1.5"), ("0", "0"), (None, None), ("   ", None)])
def test_legacy_text_normalization_removes_only_integer_float_suffix(value: object, expected: str | None) -> None:
    assert _text(value) == expected


def test_legacy_partial_rating_is_not_classified_as_unrated() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "PRE-CALIFICACIÓN", "COHERENCIA", "MALA EXP", "EMOCIÓN"])
    sheet.append(["Partial", "Artist", "=(10+9)/2", None, None, 8])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        row = preview_workbook(data.getvalue(), session)[0]
    assert row.status == "error"
    assert any("Coherence is required" in error for error in row.errors)


def test_legacy_preview_computes_grace_final_from_pre_formula_without_final_column() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "PRE-CALIFICACIÓN", "COHERENCIA", "MALA EXP", "EMOCIÓN"])
    sheet.append(["Grace", "Jeff Buckley", "=(10+10+10+10+10+10+10+10+9.5+10+10)/11", 7, 0, 10])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        preview = preview_workbook(data.getvalue(), session)
    assert preview[0].computed_pre_rating == Decimal("9.954545454545454545454545455")
    assert preview[0].computed_bad_experience == Decimal("0")
    assert float(preview[0].computed_final_rating) == pytest.approx(10.202045455)


def test_legacy_audit_bad_experience_that_cannot_be_read_is_a_warning() -> None:
    from openpyxl import Workbook

    workbook = Workbook(); sheet = workbook.active
    sheet.append(["ÁLBUM", "ARTISTA", "PRE-CALIFICACIÓN", "COHERENCIA", "MALA EXP", "EMOCIÓN"])
    sheet.append(["Promises", "Floating Points", "=(9.7+8.7+7.3+8.4+9.8+7+6.2+6.6+9.9+8+9.8+8.1)/12", 7, "not cached", 9])
    data = BytesIO(); workbook.save(data)
    with SessionLocal() as session:
        preview = preview_workbook(data.getvalue(), session)
    row = preview[0]
    assert row.status == "warning"
    assert not row.errors
    assert row.computed_pre_rating is not None
    assert row.computed_bad_experience is not None
    assert row.computed_final_rating is not None
    assert any("Aftertone will use the computed value" in warning for warning in row.warnings)


def test_revisit_mark_is_current_album_state_and_survives_new_revision() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        marked = mark_for_revisit(album.id, RevisitUpdate(reason="Revisit after release hype."), session)
    assert marked.needs_revisit is True
    assert marked.revisit_reason == "Revisit after release hype."
    assert marked.revisit_marked_at is not None
    with SessionLocal() as session:
        create_revision(album.id, RatingRevisionCreate.model_validate(revision_payload(album)), session)
        assert get_album(album.id, session).needs_revisit is True
        cleared = clear_revisit_mark(album.id, session)
    assert cleared.needs_revisit is False
    assert cleared.revisit_reason is None


def test_legacy_reconciliation_requires_explicit_unique_mapping_and_preserves_unrated_tracks() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        third_track = Track(album_id=album.id, position=3, title="Last Goodbye")
        legacy = LegacyRating(
            album_id=album.id, coherence=Decimal("7"), emotion=Decimal("8"),
            extracted_scores=["10", "9.5"], pre_formula="=(10+9.5)/2",
        )
        session.add_all([third_track, legacy])
        session.commit()
        session.refresh(legacy)
        session.refresh(third_track)

        payload = LegacyReconciliationRequest(
            track_mappings=[
                LegacyTrackMapping(legacy_score_index=0, track_id=album.tracks[0].id),
                LegacyTrackMapping(legacy_score_index=1, track_id=album.tracks[1].id),
            ]
        )
        revision = reconcile_legacy_rating(session, legacy, payload)
        session.commit()
        session.refresh(legacy)
        session.refresh(revision)

        assert legacy.reconciliation_status == "reconciled"
        assert legacy.reconciled_revision_id == revision.id
        snapshots = {item.track_id: item for item in revision.track_ratings}
        assert snapshots[album.tracks[0].id].score == Decimal("10")
        assert snapshots[album.tracks[0].id].include_in_pre_rating is True
        assert snapshots[third_track.id].score is None
        assert snapshots[third_track.id].include_in_pre_rating is False
        with pytest.raises(ReconciliationError):
            reconcile_legacy_rating(session, legacy, payload)


def test_legacy_reconciliation_rejects_duplicate_track_or_missing_score_mapping() -> None:
    album = create_test_album()
    with SessionLocal() as session:
        legacy = LegacyRating(
            album_id=album.id, coherence=Decimal("7"), emotion=Decimal("8"),
            extracted_scores=["10", "9"], pre_formula="=(10+9)/2",
        )
        session.add(legacy)
        session.commit()
        session.refresh(legacy)
        duplicate_track = LegacyReconciliationRequest(
            track_mappings=[
                LegacyTrackMapping(legacy_score_index=0, track_id=album.tracks[0].id),
                LegacyTrackMapping(legacy_score_index=1, track_id=album.tracks[0].id),
            ]
        )
        with pytest.raises(ReconciliationError, match="at most one"):
            reconcile_legacy_rating(session, legacy, duplicate_track)
        missing_score = LegacyReconciliationRequest(
            track_mappings=[LegacyTrackMapping(legacy_score_index=0, track_id=album.tracks[0].id)]
        )
        with pytest.raises(ReconciliationError, match="every legacy score"):
            reconcile_legacy_rating(session, legacy, missing_score)
        assert legacy.reconciliation_status == "pending"
