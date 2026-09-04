from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from unicodedata import normalize

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.imports.pre_formula import PreFormulaError, legacy_adjustment_value, parse_pre_formula
from app.models.music import Album, LegacyRating, ReleaseType
from app.ratings.calculator import TrackScore, calculate_rating

TOLERANCE = Decimal("0.001")


def _header(value: object) -> str:
    return "".join(character for character in normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().upper() if character.isalnum())


def _decimal(value: object, field_name: str, *, required: bool = True) -> Decimal | None:
    if value is None or value == "":
        if required:
            raise ValueError(f"{field_name} is required.")
        return None
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation as error:
        raise ValueError(f"{field_name} is not a number.") from error


def _bad_experience(value: object) -> Decimal | None:
    result = _decimal(value, "Legacy bad experience", required=False)
    if result is None:
        return None
    result = result / Decimal("100") if result > 1 else result
    if not Decimal("0") <= result <= Decimal("1"):
        raise ValueError("Legacy bad experience must be a percentage between 0 and 100.")
    return result


@dataclass
class LegacyPreviewRow:
    row_number: int
    title: str | None = None
    artist: str | None = None
    year: int | None = None
    decade: str | None = None
    genre: str | None = None
    pre_formula: str | None = None
    extracted_scores: list[Decimal] = field(default_factory=list)
    legacy_final_rating: Decimal | None = None
    legacy_pre_rating: Decimal | None = None
    legacy_bad_experience: Decimal | None = None
    coherence: Decimal | None = None
    emotion: Decimal | None = None
    computed_pre_rating: Decimal | None = None
    computed_bad_experience: Decimal | None = None
    computed_final_rating: Decimal | None = None
    status: str = "ready"
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def warn(self, message: str) -> None:
        self.warnings.append(message)
        if self.status == "ready":
            self.status = "warning"

    def error(self, message: str) -> None:
        self.errors.append(message)
        self.status = "error"


def _column_map(values: tuple[object, ...]) -> dict[str, int]:
    aliases = {"ALBUM": "title", "ARTISTA": "artist", "ANO": "year", "DECADA": "decade", "GENERO": "genre", "PRECALIFICACION": "pre", "COHERENCIA": "coherence", "MALAEXP": "bad", "EMOCION": "emotion"}
    return {aliases[key]: index for index, value in enumerate(values) if (key := _header(value)) in aliases}


def _blank(value: object) -> bool:
    return value is None or not str(value).strip()


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, Decimal) and value == value.to_integral_value():
        return str(int(value))
    text = str(value).strip()
    return text or None


def preview_workbook(data: bytes, session: Session) -> list[LegacyPreviewRow]:
    try:
        formulas = load_workbook(BytesIO(data), data_only=False, read_only=True).active
        cached = load_workbook(BytesIO(data), data_only=True, read_only=True).active
    except Exception as error:
        raise ValueError("Workbook could not be opened.") from error
    headers = next(formulas.iter_rows(min_row=1, max_row=1, values_only=True), ())
    columns = _column_map(headers)
    missing = {"title", "artist", "pre", "coherence", "emotion"} - set(columns)
    if missing:
        raise ValueError(f"Workbook is missing required headers: {', '.join(sorted(missing))}.")
    cached_rows = {index: row for index, row in enumerate(cached.iter_rows(min_row=2, values_only=True), start=2)}
    rows: list[LegacyPreviewRow] = []
    seen: set[tuple[str, str]] = set()
    for row_number, values in enumerate(formulas.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        row = LegacyPreviewRow(row_number=row_number)
        try:
            row.title = _text(values[columns["title"]])
            row.artist = _text(values[columns["artist"]])
            if not row.title or not row.artist:
                raise ValueError("Album and artist are required.")
            row.year = int(_decimal(values[columns["year"]], "Year", required=False)) if columns.get("year") is not None and values[columns["year"]] not in (None, "") else None
            row.decade = str(values[columns["decade"]]).strip() if columns.get("decade") is not None and values[columns["decade"]] else None
            row.genre = str(values[columns["genre"]]).strip() if columns.get("genre") is not None and values[columns["genre"]] else None
            rating_values = [values[columns["pre"]], values[columns["coherence"]], values[columns["emotion"]]]
            if columns.get("bad") is not None:
                rating_values.append(values[columns["bad"]])
            if all(_blank(value) for value in rating_values):
                row.status = "unrated"
                key = (row.artist.casefold(), row.title.casefold())
                if key in seen or session.scalar(select(Album.id).where(func.lower(Album.title) == row.title.casefold(), func.lower(Album.artist) == row.artist.casefold())) is not None:
                    row.warn("An album with this artist and title already exists.")
                seen.add(key)
                rows.append(row)
                continue
            row.pre_formula = str(values[columns["pre"]] or "").strip()
            row.extracted_scores, denominator, suffix = parse_pre_formula(row.pre_formula)
            row.coherence = _decimal(values[columns["coherence"]], "Coherence")
            row.emotion = _decimal(values[columns["emotion"]], "Emotion")
            result = calculate_rating([TrackScore(score) for score in row.extracted_scores], coherence=row.coherence, emotion=row.emotion)
            row.computed_pre_rating, row.computed_bad_experience, row.computed_final_rating = result.pre_rating, result.bad_experience, result.final_rating
            if denominator != len(row.extracted_scores):
                row.warn(f"Legacy PRE denominator ({denominator}) differs from extracted score count ({len(row.extracted_scores)}). Aftertone uses the extracted scores for canonical computation.")
            adjustment = legacy_adjustment_value(suffix)
            if adjustment:
                row.warn(f"Legacy PRE formula contains a manual adjustment ({adjustment:+f}). Aftertone does not apply it to canonical computation.")
            cached_value = cached_rows[row_number][columns["pre"]]
            try:
                row.legacy_pre_rating = _decimal(cached_value, "Legacy PRE", required=False)
            except ValueError:
                row.warn("Legacy PRE could not be read; formula-derived PRE will be retained for audit.")
            else:
                if row.legacy_pre_rating is None:
                    row.warn("Cached PRE value is unavailable; formula-derived PRE will be retained for audit.")
            if columns.get("bad") is not None:
                try:
                    row.legacy_bad_experience = _bad_experience(values[columns["bad"]])
                except ValueError:
                    row.warn("Legacy bad experience could not be read; Aftertone will use the computed value.")
                else:
                    if row.legacy_bad_experience is None:
                        row.warn("Legacy bad experience could not be read; Aftertone will use the computed value.")
            if row.year is None:
                row.warn("Year is missing.")
            if row.legacy_pre_rating is not None and abs(row.legacy_pre_rating - row.computed_pre_rating) > TOLERANCE:
                row.warn("Legacy PRE differs from the computed PRE.")
            if row.legacy_bad_experience is not None and abs(row.legacy_bad_experience - row.computed_bad_experience) > TOLERANCE:
                row.warn("Legacy bad experience differs from the computed value.")
            key = (row.artist.casefold(), row.title.casefold())
            if key in seen or session.scalar(select(Album.id).where(func.lower(Album.title) == row.title.casefold(), func.lower(Album.artist) == row.artist.casefold())) is not None:
                row.warn("An album with this artist and title already exists.")
            seen.add(key)
        except (ValueError, PreFormulaError) as error:
            row.error(str(error))
        rows.append(row)
    return rows


def commit_rows(rows: list[LegacyPreviewRow], selected_row_numbers: set[int], session: Session) -> tuple[int, int, int]:
    imported = skipped = failed = 0
    for row in rows:
        if row.row_number not in selected_row_numbers or row.status == "error":
            skipped += 1
            continue
        try:
            if not row.title or not row.artist:
                raise ValueError("Row is incomplete.")
            if session.scalar(select(Album.id).where(func.lower(Album.title) == row.title.casefold(), func.lower(Album.artist) == row.artist.casefold())) is not None:
                skipped += 1
                continue
            album = Album(title=row.title, artist=row.artist, release_year=row.year, release_type=ReleaseType.ALBUM)
            if row.status == "unrated":
                session.add(album); session.commit(); imported += 1
                continue
            if row.coherence is None or row.emotion is None or not row.pre_formula:
                raise ValueError("Row is incomplete.")
            legacy = LegacyRating(album=album, legacy_final_rating=row.legacy_final_rating, legacy_pre_rating=row.legacy_pre_rating, legacy_bad_experience=row.legacy_bad_experience, coherence=row.coherence, emotion=row.emotion, extracted_scores=[str(score) for score in row.extracted_scores], pre_formula=row.pre_formula, computed_pre_rating=row.computed_pre_rating, computed_bad_experience=row.computed_bad_experience, computed_final_rating=row.computed_final_rating, legacy_decade=row.decade, legacy_genre=row.genre)
            session.add_all([album, legacy]); session.commit(); imported += 1
        except Exception:
            session.rollback(); failed += 1
    return imported, skipped, failed
