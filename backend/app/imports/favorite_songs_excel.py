from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from unicodedata import normalize
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.favorite_songs.calculator import calculate_favorite_song_score
from app.models.music import Album, AlbumArtist, FavoriteSongEntry, Track, TrackArtist, TrackArtistRole


def key(value: object) -> str:
    return " ".join(normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().casefold().split())


def title_key(value: object) -> str:
    """Match legacy titles without treating structured featured credits as title identity."""
    visible = key(value)
    visible = re.sub(r"\s*[\(\[]\s*(?:feat\.?|featuring|ft\.?)\s+.*?[\)\]]\s*$", "", visible)
    return "".join(character for character in visible if character.isalnum())


def header_key(value: object) -> str:
    return "".join(character for character in normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().casefold() if character.isalnum())


def decimal(value: object, name: str, maximum: Decimal) -> Decimal:
    if value is None or str(value).strip() == "": raise ValueError(f"{name} is required.")
    try: result = Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation as error: raise ValueError(f"{name} is not a number.") from error
    if not Decimal("0") <= result <= maximum: raise ValueError(f"{name} must be between 0 and {maximum}.")
    return result


@dataclass
class FavoriteImportRow:
    source_row: int; source_position: str | None; song: str | None; artist: str | None; album: str | None; year: int | None; genre: str | None
    base_score: Decimal | None = None; emotional_connection: Decimal | None = None; replay_value: Decimal | None = None; historical_relevance: Decimal | None = None; originality: Decimal | None = None; legacy_final_score: Decimal | None = None; calculated_final_score: Decimal | None = None
    status: str = "invalid"; matched_track_id: int | None = None; candidates: list[dict] = field(default_factory=list); warnings: list[str] = field(default_factory=list); errors: list[str] = field(default_factory=list)


ALIASES = {"posicion":"position", "cancion":"song", "artista":"artist", "album":"album", "ano":"year", "genero":"genre", "puntuacion":"base", "conexionemocional":"emotional", "replayvalue":"replay", "releimphistorico":"historical", "originalidad":"originality", "calificacion":"final"}

def preview_favorite_workbook(data: bytes, session: Session) -> list[FavoriteImportRow]:
    try: sheet = load_workbook(BytesIO(data), data_only=True, read_only=True)["100 songs"]
    except KeyError as error: raise ValueError('Workbook must contain a worksheet named "100 songs".') from error
    except Exception as error: raise ValueError("Workbook could not be opened.") from error
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    columns = {ALIASES.get(header_key(value)): index for index, value in enumerate(headers) if ALIASES.get(header_key(value))}
    required = {"song", "artist", "album", "base", "emotional", "replay", "historical", "originality"}
    if missing := required - set(columns): raise ValueError(f"Workbook is missing required headers: {', '.join(sorted(missing))}.")
    tracks = session.scalars(select(Track).options(selectinload(Track.album).selectinload(Album.artist_credits).selectinload(AlbumArtist.artist), selectinload(Track.artist_credits).selectinload(TrackArtist.artist), selectinload(Track.favorite_song_entry))).all()
    seen: set[int] = set(); rows = []
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value is not None and str(value).strip() for value in values): continue
        row = FavoriteImportRow(number, str(values[columns["position"]]).strip() if "position" in columns and values[columns["position"]] is not None else None, str(values[columns["song"]]).strip() or None, str(values[columns["artist"]]).strip() or None, str(values[columns["album"]]).strip() or None, int(values[columns["year"]]) if "year" in columns and values[columns["year"]] is not None else None, str(values[columns["genre"]]).strip() or None if "genre" in columns and values[columns["genre"]] is not None else None)
        try:
            row.base_score=decimal(values[columns["base"]], "Puntuación", Decimal("10")); row.emotional_connection=decimal(values[columns["emotional"]], "Conexión emocional", Decimal("5")); row.replay_value=decimal(values[columns["replay"]], "Replay value", Decimal("5")); row.historical_relevance=decimal(values[columns["historical"]], "Rel. e Imp. Histórico", Decimal("5")); row.originality=decimal(values[columns["originality"]], "Originalidad", Decimal("5"))
            row.calculated_final_score=calculate_favorite_song_score(base_score=row.base_score, emotional_connection=row.emotional_connection, replay_value=row.replay_value, historical_relevance=row.historical_relevance, originality=row.originality)
            if "final" in columns and values[columns["final"]] is not None:
                try: row.legacy_final_score=Decimal(str(values[columns["final"]]).replace(",", "."))
                except InvalidOperation: row.warnings.append("Excel rating is unavailable.")
                if row.legacy_final_score is not None and abs(row.legacy_final_score-row.calculated_final_score)>Decimal("0.01"): row.warnings.append("Excel rating differs from Aftertone rating.")
            matches=[]; title_matches=[]; title_album_matches=[]
            for track in tracks:
                artists=[credit.artist.name for credit in track.artist_credits if credit.role == TrackArtistRole.PRIMARY] or [credit.artist.name for credit in track.album.artist_credits]
                if title_key(track.title) != title_key(row.song):
                    continue
                title_matches.append((track, artists))
                if key(track.album.title) == key(row.album):
                    title_album_matches.append((track, artists))
                    if any(key(name)==key(row.artist) for name in artists): matches.append(track)
            nearby = matches or [track for track, _ in title_album_matches] or [track for track, _ in title_matches]
            row.candidates=[{"track_id": track.id, "title": track.title, "album": track.album.title, "year": track.album.release_year, "artists": ", ".join(credit.artist.name for credit in track.artist_credits if credit.role == TrackArtistRole.PRIMARY) or ", ".join(credit.artist.name for credit in track.album.artist_credits)} for track in nearby]
            if len(matches)==1:
                track=matches[0]; row.matched_track_id=track.id
                if track.id in seen: row.status="invalid"; row.errors.append("Duplicate track in workbook.")
                elif track.favorite_song_entry: row.status="already_ranked"
                else: row.status="ready"; seen.add(track.id)
            elif len(matches)>1: row.status="review"; row.warnings.append("More than one track has the same song, album, and artist identity. Choose the intended version.")
            else:
                row.status="not_found"
                if not title_matches:
                    row.warnings.append(f'No existing track title matches “{row.song}”.')
                elif not title_album_matches:
                    row.warnings.append(f'Found the song title, but no matching album for “{row.album}”. Nearby matches: ' + "; ".join(f"{track.title} — {track.album.title}" for track, _ in title_matches[:3]))
                else:
                    row.warnings.append(f'Found matching song and album, but artist credits do not match “{row.artist}”. Existing credits: ' + "; ".join(", ".join(artists) for _, artists in title_album_matches[:3]))
        except ValueError as error: row.errors.append(str(error)); row.status="invalid"
        rows.append(row)
    return rows
